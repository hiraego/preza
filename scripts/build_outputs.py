"""Сборка итоговых файлов из build/jpg/*.jpg:
  dist/...pdf  — каждая страница = одна картинка (никакого живого текста)
  dist/...pptx — каждый слайд = картинка на весь слайд + заметки докладчика
"""
import json
import pathlib
import sys

import img2pdf
from pptx import Presentation
from pptx.util import Emu

ROOT = pathlib.Path(__file__).resolve().parent.parent
JPG = sorted((ROOT / 'build/jpg').glob('*.jpg'))
NOTES = json.loads((ROOT / 'build/notes.json').read_text(encoding='utf-8'))
DIST = ROOT / 'dist'
DIST.mkdir(exist_ok=True)
NAME = 'Аладько_Известные_зарубежные_экологи'

if len(JPG) != len(NOTES):
    sys.exit(f'слайдов {len(JPG)}, заметок {len(NOTES)}: сначала node scripts/render.mjs')

# PDF 16:9, 338.67 x 190.5 мм (как слайд PowerPoint)
W_PT, H_PT = img2pdf.in_to_pt(13.333), img2pdf.in_to_pt(7.5)
layout = img2pdf.get_layout_fun((W_PT, H_PT), fit=img2pdf.FitMode.into)
pdf = DIST / f'{NAME}.pdf'
pdf.write_bytes(img2pdf.convert([str(p) for p in JPG], layout_fun=layout,
                                title='Известные зарубежные экологи', author='Аладько Иван Игоревич'))

# PPTX
prs = Presentation()
prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)
blank = prs.slide_layouts[6]
for img, note in zip(JPG, NOTES):
    slide = prs.slides.add_slide(blank)
    slide.shapes.add_picture(str(img), 0, 0, width=prs.slide_width, height=prs.slide_height)
    slide.notes_slide.notes_text_frame.text = note
prs.core_properties.title = 'Известные зарубежные экологи'
prs.core_properties.author = 'Аладько Иван Игоревич'
pptx = DIST / f'{NAME}.pptx'
prs.save(pptx)

words = sum(len(n.split()) for n in NOTES)
print(f'{len(JPG)} слайдов · заметки {words} слов')
for f in (pdf, pptx):
    print(f'{f.relative_to(ROOT)}  {f.stat().st_size / 1e6:.1f} МБ')
