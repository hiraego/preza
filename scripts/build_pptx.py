"""Нативный PPTX из build/pptx/plan.json (его готовит scripts/export_pptx.mjs).

Каждый слайд: фон заливкой, графика каждого блока отдельной картинкой (с прозрачностью),
весь текст настоящими текстовыми полями (шрифт, размер, цвет, межбуквенный интервал, переносы как в браузере),
заметки докладчика. Шрифты встраиваются в файл (EOT .fntdata), поэтому на чужом компьютере не должны «поплыть».
"""
import io
import json
import pathlib
import re
import struct
import sys

from fontTools.ttLib import TTFont
from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.opc.package import Part
from pptx.opc.packuri import PackURI
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

ROOT = pathlib.Path(__file__).resolve().parent.parent
SRC = ROOT / 'build/pptx'
FONTS = ROOT / 'assets/fonts'
OUT = ROOT / 'dist/Аладько_Известные_зарубежные_экологи.pptx'
PX = 6350  # EMU в одном CSS-пикселе (1920 px = 13.333 дюйма, 1 px = 0.5 pt)

# (семейство CSS, вес, курсив) -> (имя гарнитуры в PPTX, жирный, курсив, файл)
FACES = {
    ('Inter Tight', 400, False): ('Inter Tight', False, False, 'InterTight_400Regular.ttf'),
    ('Inter Tight', 500, False): ('Inter Tight Medium', False, False, 'InterTight_500Medium.ttf'),
    ('Inter Tight', 600, False): ('Inter Tight SemiBold', False, False, 'InterTight_600SemiBold.ttf'),
    ('Inter Tight', 800, False): ('Inter Tight ExtraBold', False, False, 'InterTight_800ExtraBold.ttf'),
    ('Spectral', 400, True): ('Spectral', False, True, 'Spectral_400Regular_Italic.ttf'),
    ('Spectral', 500, True): ('Spectral Medium', False, True, 'Spectral_500Medium_Italic.ttf'),
    ('IBM Plex Mono', 400, False): ('IBM Plex Mono', False, False, 'IBMPlexMono_400Regular.ttf'),
    ('IBM Plex Mono', 500, False): ('IBM Plex Mono Medium', False, False, 'IBMPlexMono_500Medium.ttf'),
}


def face(style):
    fam = style['family'] if style['family'] in ('Inter Tight', 'Spectral', 'IBM Plex Mono') else 'Inter Tight'
    it = style['italic'] if fam == 'Spectral' else False
    if fam == 'Spectral':
        it = True
    weights = sorted({w for (f, w, i) in FACES if f == fam and i == it})
    w = min(weights, key=lambda v: (abs(v - style['weight']), -v))
    return FACES[(fam, w, it)]


def rgba(css):
    m = re.match(r'rgba?\(([^)]+)\)', css)
    v = [float(x) for x in m.group(1).replace('/', ',').split(',')]
    return RGBColor(int(v[0]), int(v[1]), int(v[2])), (v[3] if len(v) > 3 else 1.0)


# ---------- EOT для встраивания шрифтов ----------
def make_eot(path):
    data = path.read_bytes()
    t = TTFont(path)
    os2, head, name = t['OS/2'], t['head'], t['name']

    def nm(i):
        return (name.getDebugName(i) or '').encode('utf-16-le')
    pan = os2.panose
    panose = bytes([pan.bFamilyType, pan.bSerifStyle, pan.bWeight, pan.bProportion, pan.bContrast,
                    pan.bStrokeVariation, pan.bArmStyle, pan.bLetterForm, pan.bMidline, pan.bXHeight])
    italic = 1 if (os2.fsSelection & 1) else 0
    hdr = struct.pack('<LLLL', 0, len(data), 0x00020001, 0) + panose + struct.pack('<BBLHH', 1, italic, os2.usWeightClass, os2.fsType, 0x504C)
    hdr += struct.pack('<LLLL', os2.ulUnicodeRange1, os2.ulUnicodeRange2, os2.ulUnicodeRange3, os2.ulUnicodeRange4)
    hdr += struct.pack('<LL', os2.ulCodePageRange1, os2.ulCodePageRange2)
    hdr += struct.pack('<L', head.checkSumAdjustment) + struct.pack('<LLLL', 0, 0, 0, 0) + struct.pack('<H', 0)
    for i in (1, 2, 5, 4):
        s = nm(i)
        hdr += struct.pack('<H', len(s)) + s + struct.pack('<H', 0)
    hdr += struct.pack('<H', 0)  # RootStringSize = 0
    total = len(hdr) + len(data)
    return struct.pack('<L', total) + hdr[4:] + data


def embed_fonts(prs):
    RT = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/font'
    pres = prs.part._element
    lst = etree.SubElement(pres, qn('p:embeddedFontLst'))
    pres.remove(lst)
    pres.find(qn('p:notesSz')).addnext(lst)
    for k, (tf, b, i, fn) in enumerate(sorted(set(FACES.values()))):
        part = Part(PackURI(f'/ppt/fonts/font{k + 1}.fntdata'), 'application/x-fontdata', prs.part.package, make_eot(FONTS / fn))
        rid = prs.part.relate_to(part, RT)
        ef = etree.SubElement(lst, qn('p:embeddedFont'))
        etree.SubElement(ef, qn('p:font'), typeface=tf, pitchFamily='49' if 'Mono' in tf else '2', charset='0')
        slot = etree.SubElement(ef, qn('p:italic' if i else 'p:regular'))
        slot.set(qn('r:id'), rid)
    pres.set('embedTrueTypeFonts', '1')


# ---------- картинки слоев ----------
def picture(slide, png, name, bg):
    im = Image.open(SRC / png).convert('RGBA')
    a = im.getchannel('A')
    box = a.getbbox()
    if not box:
        return
    im = im.crop(box)
    a = im.getchannel('A')
    k = 1 / plan['scale']
    buf = io.BytesIO()
    # непрозрачная картинка (полупрозрачны только края от дробных координат) -> JPEG
    inner = a.crop((2, 2, max(3, im.width - 2), max(3, im.height - 2)))
    if inner.getextrema()[0] == 255:
        flat = Image.new('RGB', im.size, bg)
        flat.paste(im, mask=a)
        flat.save(buf, 'JPEG', quality=88, optimize=True, progressive=True)
    else:
        im.save(buf, 'PNG', optimize=True)
    buf.seek(0)
    pic = slide.shapes.add_picture(buf, Emu(round(box[0] * k * PX)), Emu(round(box[1] * k * PX)),
                                   Emu(round(im.width * k * PX)), Emu(round(im.height * k * PX)))
    pic.name = name


# ---------- текст ----------
def textbox(slide, t):
    lines = t['lines']
    first = lines[0]
    s0 = max(r['style']['size'] for r in first['runs'])
    # шаг строк берем из браузера и задаем точным значением: так он одинаков в PowerPoint и LibreOffice
    if len(lines) > 1:
        pitch = (lines[-1]['base'] - first['base']) / (len(lines) - 1)
    else:
        pitch = 1.2 * s0
    top = first['base'] - 0.8 * pitch
    x0 = min(L['x0'] for L in lines)
    x1 = max(L['x1'] for L in lines)
    slack = 0.04 * (x1 - x0) + 6
    if t['align'] == 'center':
        cx = t['box']['x'] + t['box']['w'] / 2
        left, width = cx - (x1 - x0) / 2 - slack, (x1 - x0) + 2 * slack
    elif t['align'] == 'right':
        left, width = x0 - slack, (x1 - x0) + slack
    else:
        left, width = x0, (x1 - x0) + slack
    height = len(lines) * pitch
    sh = slide.shapes.add_textbox(Emu(round(left * PX)), Emu(round(top * PX)), Emu(round(width * PX)), Emu(round(height * PX)))
    txt = ' '.join(''.join(r['text'] for r in L['runs']) for L in lines)
    sh.name = 'Текст: ' + txt[:40]
    tfm = sh.text_frame
    tfm.word_wrap = False
    tfm.auto_size = MSO_AUTO_SIZE.NONE
    tfm.vertical_anchor = MSO_ANCHOR.TOP
    tfm.margin_left = tfm.margin_right = tfm.margin_top = tfm.margin_bottom = 0
    para = tfm.paragraphs[0]
    para.alignment = {'center': PP_ALIGN.CENTER, 'right': PP_ALIGN.RIGHT}.get(t['align'], PP_ALIGN.LEFT)
    para.line_spacing = Pt(round(pitch * 0.5, 2))
    for li, L in enumerate(lines):
        if li:
            br = etree.SubElement(para._p, qn('a:br'))
            style_rpr(etree.SubElement(br, qn('a:rPr')), L['runs'][0]['style'])
        for r in L['runs']:
            run = para.add_run()
            run.text = r['text']
            style_rpr(run._r.get_or_add_rPr(), r['style'])


def style_rpr(rpr, st):
    tf, b, i, _ = face(st)
    rpr.set('lang', 'ru-RU')
    rpr.set('sz', str(round(st['size'] * 50)))
    rpr.set('b', '1' if b else '0')
    rpr.set('i', '1' if i else '0')
    if st['ls']:
        rpr.set('spc', str(round(st['ls'] * 50)))
    color, alpha = rgba(st['color'])
    fill = etree.SubElement(rpr, qn('a:solidFill'))
    clr = etree.SubElement(fill, qn('a:srgbClr'), val=str(color))
    if alpha < 0.999:
        etree.SubElement(clr, qn('a:alpha'), val=str(round(alpha * 100000)))
    for tag in ('a:latin', 'a:ea', 'a:cs'):
        etree.SubElement(rpr, qn(tag), typeface=tf)


plan = json.loads((SRC / 'plan.json').read_text(encoding='utf-8'))
if any(s is None for s in plan['slides']):
    sys.exit('plan.json собран не для всех слайдов: запустите node scripts/export_pptx.mjs без номеров')

prs = Presentation()
prs.slide_width, prs.slide_height = Emu(1920 * PX), Emu(1080 * PX)
blank = prs.slide_layouts[6]
for n, S in enumerate(plan['slides'], 1):
    slide = prs.slides.add_slide(blank)
    color, _ = rgba(S['bg'])
    bg = tuple(int(str(color)[i:i + 2], 16) for i in (0, 2, 4))
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color
    if S.get('bgImage'):
        picture(slide, S['bgImage'], 'Фон', bg)
    for L in S['layers']:
        if L.get('image'):
            picture(slide, L['image'], f'Графика {L["j"] + 1}', bg)
        for t in L['texts']:
            textbox(slide, t)
    slide.notes_slide.notes_text_frame.text = S['notes']
prs.core_properties.title = 'Известные зарубежные экологи'
prs.core_properties.author = 'Аладько Иван Игоревич'
embed_fonts(prs)
OUT.parent.mkdir(exist_ok=True)
prs.save(OUT)
print(f'{OUT.relative_to(ROOT)}  {OUT.stat().st_size / 1e6:.1f} МБ, слайдов {len(plan["slides"])}')
